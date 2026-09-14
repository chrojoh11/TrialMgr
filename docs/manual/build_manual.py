from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "SDDA-TrialDesk-Secretary-Manual.docx"
SHOTS = ROOT / "screenshots"

GREEN = "294F73"; GOLD = "526B83"; PALE = "F1F4F7"; LIGHT_GREEN = "E8EFF5"
INK = "18232F"; MUTED = "637080"; WHITE = "FFFFFF"; RED = "9B1C1C"; AMBER = "F1F4F7"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(.75); sec.bottom_margin = Inches(.72)
sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)
sec.header_distance = Inches(.35); sec.footer_distance = Inches(.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.16
for name, size, color, before, after in [
    ("Title", 28, GREEN, 0, 10), ("Subtitle", 13, MUTED, 0, 10),
    ("Heading 1", 18, GREEN, 15, 7), ("Heading 2", 14, GREEN, 11, 5),
    ("Heading 3", 11.5, GOLD, 8, 3),
]:
    s = styles[name]; s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = name != "Subtitle"
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for tag, val in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = tcMar.find(qn(f"w:{tag}"))
        if node is None: node=OxmlElement(f"w:{tag}"); tcMar.append(node)
        node.set(qn("w:w"), str(val)); node.set(qn("w:type"), "dxa")

def set_cell_text(cell, text, bold=False, color=INK, size=9.5):
    cell.text = ""; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(text); r.bold=bold; r.font.name="Calibri"; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)

def table(rows, widths=None, header=True):
    t=doc.add_table(rows=0, cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for ci,value in enumerate(row):
            if widths: cells[ci].width=Inches(widths[ci])
            set_cell_text(cells[ci], str(value), bold=(header and ri==0), color=WHITE if header and ri==0 else INK)
            if header and ri==0: shade(cells[ci], GREEN)
            elif ri%2==0: shade(cells[ci], "F7F8F4")
    return t

def heading(text, level=1): doc.add_heading(text, level=level)

def para(text="", bold_lead=None, italic=False):
    p=doc.add_paragraph();
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold=True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    if italic:
        for r in p.runs: r.italic=True
    return p

def bullet(text, level=0):
    p=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2"); p.add_run(text); return p

def step(number, title, detail):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.05); p.paragraph_format.first_line_indent=Inches(-.05)
    r=p.add_run(f"{number}. {title}. "); r.bold=True; r.font.color.rgb=RGBColor.from_string(GREEN)
    p.add_run(detail); return p

def callout(label, text, kind="note"):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    c=t.cell(0,0); c.width=Inches(6.8); shade(c, AMBER if kind=="warning" else LIGHT_GREEN)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(label.upper()+": "); r.bold=True; r.font.color.rgb=RGBColor.from_string(RED if kind=="warning" else GREEN)
    p.add_run(text); margins(c,120,160,120,160)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def image(name, caption, width=6.8):
    path=SHOTS/name
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path), width=Inches(width))
    c=doc.add_paragraph(caption); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after=Pt(8)
    for r in c.runs: r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(MUTED)

def page_break(): doc.add_page_break()

def hyperlink(paragraph, text, url):
    part=paragraph.part; rid=part.relate_to(url,"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
    h=OxmlElement("w:hyperlink"); h.set(qn("r:id"),rid); run=OxmlElement("w:r"); rPr=OxmlElement("w:rPr")
    color=OxmlElement("w:color"); color.set(qn("w:val"),GREEN); rPr.append(color)
    u=OxmlElement("w:u"); u.set(qn("w:val"),"single"); rPr.append(u); run.append(rPr)
    txt=OxmlElement("w:t"); txt.text=text; run.append(txt); h.append(run); paragraph._p.append(h)

# Header/footer
hp=sec.header.paragraphs[0]; hp.text="SDDA TrialDesk | Secretary Operations Manual"; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.add_run("SDDA TrialDesk  •  July 2026 Rulebook workflow  •  ")
fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); fp._p.append(fld)
for r in fp.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)

# Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(72); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("SDDA"); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(GOLD)
p=doc.add_paragraph("TrialDesk Secretary Manual", style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph("Create, prepare, run, score, place, export, and close an SDDA trial", style="Subtitle"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph().paragraph_format.space_after=Pt(20)
callout("Purpose", "A practical, step-by-step operating manual for trial secretaries using SDDA TrialDesk. It follows Sporting Detection Dogs Association Master Rule Book v5.1, effective July 1, 2026, and the implemented TrialDesk workflow.")
table([
    ["Manual version", "Rules authority", "Application scope"],
    ["1.4 • September 14, 2026", "SDDA Master Rule Book v5.1", "Scent, Games, and Combined trials"],
], widths=[1.8,2.25,2.75])
para("This manual explains the secretary’s administrative workflow. The current SDDA rulebook, sanction approval, premium list, judge’s instructions, and official SDDA forms remain controlling if they conflict with this guide.", italic=True)
para("Updated in this edition: every primary TrialDesk page was re-audited, current workflow screenshots were added, Trial Team and Activity Journal instructions were expanded, and a page-and-button glossary was added. Screenshots were captured from Example Trial without saving operational changes.", bold_lead="Updated in this edition:")

page_break(); heading("Contents and operating sequence",1)
for item in [
    "1. Before TrialDesk: sanctioning, premium list, venue, judge, and supplies",
    "2. Create the draft trial and select its format and dates",
    "3. Configure offerings, FEO, judges, fees, and public details",
    "4. Open entries, share the entry form, import CSV responses, and accept entries",
    "5. Build and reorganize the running order",
    "6. Apply component-specific move-ups",
    "7. Generate and verify official score sheets",
    "8. Trial-day desk workflow",
    "9. Enter scores and make audited corrections",
    "10. Review placements and Title Watch",
    "11. Finances, official workbook, backup, and closeout",
    "12. Troubleshooting and final checklists",
    "13. Page and button glossary",
]: bullet(item)
callout("Recommended practice", "Complete a full mock trial before using TrialDesk at a live event. Confirm the entry form, one score sheet from every offered class, the running-order export, scoring, placements, workbook, backup, and printing on the equipment that will be used at the trial.")

page_break(); heading("1. Before opening TrialDesk",1)
heading("1.1 Obtain SDDA approval",2)
step(1,"Confirm the judge and venue","Consult the judge about whether the search areas are suitable for every proposed level, component, and Game.")
step(2,"Submit the SDDA trial application","The rulebook requires the relevant application, judge, venue, and contact information at least 60 days before the trial.")
step(3,"Wait for publication","After approval and SDDA publication, prepare the Premium List and entry form. Do not distribute the public entry form earlier than allowed by the current SDDA instructions.")
step(4,"Document changes","If the judge, venue, offering, schedule, or Premium changes, notify entrants and SDDA promptly and be prepared to refund as required.")
heading("1.2 Gather setup information",2)
table([
    ["Information", "What to have ready"],
    ["Trial identity", "Trial name, host club, venue name, full address, and trial dates"],
    ["Sanction details", "SDDA trial number for each day when issued; judge name and substitutions"],
    ["Offerings", "Scent levels/components and/or Games; any FEO limits; capacities"],
    ["Pricing", "Scent component/three-component/Elite fees; separate Regular and FEO Games fees"],
    ["Competitor information", "Secretary contact, payment instructions, cancellation/refund terms"],
    ["Operations", "Search-area sequence, officials, timers, stewards, staging, printing, backup plan"],
], widths=[1.45,5.35])
heading("1.3 Physical host checklist",2)
for x in ["Prepared scent kit and separately stored clean supplies","Correct official score sheets, clipboards, pens, calculators, stopwatches, and spares","Safe search areas, clearly marked boundaries and starting marks","Containers, vehicles/objects, distractions, replacement hide objects","Enough judges, timers, stewards, gate staff, and a sight-screened staging area","A posted running order and a check-in station","Printer access or all required judge packets printed in advance"]: bullet(x)

page_break(); heading("2. Create a trial",1)
image("01-dashboard.png","Figure 1. TrialDesk dashboard. Select an existing trial or start a new draft; the operation links follow the trial through setup, entries, running order, sheets, scoring, results, export, and closeout.",6.8)
step(1,"Sign in","Use the secretary account assigned to the trial.")
step(2,"Select Create trial","Use the left menu or the dashboard button.")
step(3,"Enter trial details","Use a recognizable trial name, the legal host/club name, and the complete venue address shown to competitors.")
step(4,"Choose the trial format","Choose Scent, Games, or Combined. Combined presents both workflows in the same trial workspace.")
step(5,"Add the trial dates","Enter unique dates in chronological order. Each becomes a separate trial day in TrialDesk. TrialDesk currently accepts one to four dates; this is an application limit, not an SDDA rule.")
step(6,"Create Draft Trial","A draft does not accept public entries. It can be completed as information arrives.")
image("02-create-trial.png","Figure 2. Initial trial creation. Trial numbers, judges, offerings, fees, and public instructions are completed on the next screen and may be amended later.",6.8)
callout("Rule distinction", "SDDA defines a trial as a single offering of any level within a 12-hour period with a unique trial number. A consecutive set of trials is a trial event. Enter the SDDA numbers exactly as issued for each TrialDesk day.","warning")

heading("2.1 Edit an existing trial",2)
step(1,"Open All Trials","Find the trial and select Edit trial. The same editor is available from the trial's edit address.")
step(2,"Update trial identity","Change the trial name, host club, venue/address, or Scent/Games/Combined format, then select Save trial changes.")
step(3,"Update trial days","Revise dates, add days up to the four-day TrialDesk limit, or remove only the final day. Enter dates chronologically and select Save trial days.")
para("Date changes preserve each existing Day number and everything attached to it. Removing a day is blocked until its offerings are cleared and it contains no competitor runs. Reopen a completed trial before making schedule changes.")

heading("2.2 Copy a trial",2)
step(1,"Select Copy trial","Use the control on an All Trials card or at the bottom of Edit trial.")
step(2,"Name the new trial and enter its dates","Provide one new chronological date for every day in the source trial.")
step(3,"Review the copy boundary","TrialDesk copies the club, venue, format, public instructions, pricing, capacities, FEO settings, and offerings. It does not copy entries, dogs, runs, scores, payments, expenses, trial numbers, or judges.")
step(4,"Create draft copy","Open the new draft, review every setting, enter new trial numbers and judges, and only then open entries.")

page_break(); heading("3. Configure the draft trial",1)
image("11-trial-setup.png","Figure 3. Trial setup dashboard. The workflow row, operational summary, setup checklist, and trial operation links keep the secretary’s work in one place.",6.8)
heading("3.1 Use the setup checklist",2)
para("Open the trial. The setup screen tracks offerings, pricing, competitor-facing details, trial numbers, judges, and whether the entry form is open. Amber items may remain pending while the trial is a draft, but resolve them before final packets and exports.")
heading("3.2 Enter public details",2)
step(1,"Enter the trial secretary","Provide the name, monitored email address, and optional phone number.")
step(2,"Write payment instructions","State when payment is due, method, recipient, reference information, and the consequence of non-payment.")
step(3,"Enter cancellation/refund terms","Match the approved Premium List. Keep this text current if SDDA-approved arrangements change.")
step(4,"Save public details","These details appear on the competitor entry form.")
heading("3.3 Configure Scent offerings",2)
for x in [
    "For each day, select the offered level: Started, Advanced, Excellent, or Elite.",
    "Within each level, select Container, Interior, and/or Exterior. Elite must offer all three components. If Started Container is offered, it must be held first and indoors.",
    "Do not choose Amateur or Working during trial creation. TrialDesk makes both streams available per selected component on the entry form; Elite has no stream.",
    "Enter a component judge when it differs from the day judge. Both Amateur and Working use that component judge.",
    "Select Allow FEO entries only where the organizer permits FEO. The public FEO checkbox appears only for those offerings.",
]: bullet(x)
heading("3.4 Configure Games",2)
for x in [
    "Select Aerial, Distance, Speed, and/or Team for each day.",
    "Enter the judge, capacity, Regular fee, and—only when permitted—enable FEO and enter its fee.",
    "Aerial entrants select High or Highfly. Team entrants identify a requested partner; the secretary finalizes pairings and order.",
    "Remember the rulebook-specific second-dog provisions: Aerial and Distance allow second dogs; Speed and Team do not.",
]: bullet(x)
heading("3.5 Enter day details and fees",2)
step(1,"Enter SDDA trial number and day judge","They may be saved or replaced later if the assignment changes.")
step(2,"Enter Scent pricing","Set individual component, three-component package, and Elite fees. TrialDesk uses them to calculate accepted-entry balances.")
step(3,"Save setup","Do this before opening or copying the entry form link.")
callout("FEO", "FEO is subordinate to Regular entries. SDDA permits FEO only after regular entries are accepted and allows the organizer to limit FEO by component or level when those limits are stated on the entry form.")
heading("3.6 Give another secretary access",2)
image("12-trial-team.png","Figure 4. Trial Team. Add an already registered TrialDesk user as Secretary, Assistant, or Viewer.",6.4)
step(1,"Open Trial Team","Use the trial sidebar or the Trial operations buttons on Trial Details.")
step(2,"Enter the registered email","The person must already have registered and confirmed that exact email address.")
step(3,"Choose a role","Secretary has operational access, Assistant supports trial work, and Viewer has read-only access.")
step(4,"Select Add to trial","Re-enter an existing member’s email to change the role. Use the trash button only to remove a non-owner member.")

page_break(); heading("4. Entries: collect, import, review, and accept",1)
image("13-entry-roster.png","Figure 5. Entry roster. Summary totals, search, status filter, mailing-list export, CSV import, entry editing, and secretary decisions are available from one page.",6.8)
heading("4.1 Open and share the built-in entry form",2)
step(1,"Review the preview","Confirm dates, offerings, streams, FEO availability, prices, contact information, payment instructions, declaration, and cancellation terms.")
step(2,"Open entries","Use the trial’s entry-status control only after the public form is ready.")
step(3,"Copy entry form link","Send the public link or place it in the approved Premium/club communication.")
para("Competitors enter handler/contact details, participant number, dog call name, SDDA registration number or registration pending, breed, per-component Amateur/Working stream, offered FEO selections, formal alerts, title note, and reactivity. The separate Registered Name question and competitor-reported Gold counts are intentionally not used.")
heading("4.2 Import Google Form CSV responses",2)
step(1,"Export the Google Form responses as CSV","Keep the header row intact.")
step(2,"Open Entries / Entry roster","Choose the import control and select the CSV.")
step(3,"Read every warning","TrialDesk recognizes familiar Google headings, expands all-component selections into individual runs, and rejects selections not offered for that day/level/stream.")
step(4,"Resolve—not ignore—row errors","Correct missing identity fields, stream inconsistencies, duplicate submissions, invalid dates, or unavailable offerings before relying on the roster.")
image("04-form-responses-headers.png","Figure 6. Sanitized Google Form response headings from the original SDDA workflow. TrialDesk imports the response selections into the unified entry roster.",6.8)
heading("4.3 Review and accept entries",2)
for x in [
    "Open each received entry and verify contact information, call name, SDDA number/pending status, breed, stream, formal alerts, reactivity, and requested runs.",
    "Confirm capacity and payment handling before changing Received to Accepted.",
    "Waitlisted, withdrawn, and cancelled entries must not appear in final running orders, score packets, scoring, or fee expectations.",
    "The secretary can edit an entry after acceptance when a correction is necessary; the activity journal records material changes.",
    "Use the private receipt/edit link for competitor amendments while entries remain open. Do not publish private edit tokens.",
]: bullet(x)
callout("Check-in requirement", "On trial day, confirm the dog’s SDDA registration number and current running order before the handler briefing. Resolve registration-pending entries before final SDDA submission.","warning")

page_break(); heading("5. Build and reorganize the running order",1)
image("14-running-order.png","Figure 7. Running Orders. Select a day, class, and component; apply SDDA order, assign run groups, approve move-ups, reposition runs, and save before exporting.",6.8)
heading("5.1 Generate the initial order",2)
step(1,"Accept entries first","Only accepted entries become operational runs.")
step(2,"Open Running orders","Select the trial day and generate or review the order for each level/component and Game.")
step(3,"Apply SDDA group order","TrialDesk orders Scent runs as Officials, Regular, Second dog, FEO, then BIS. Games show Regular before FEO.")
step(4,"Review conflicts","Check handlers with multiple dogs, officials who must run before duties, reactive-dog notes, Team partners, simultaneous rings, and travel between search areas.")
heading("5.2 Change group or position",2)
for x in [
    "Use the run-group control when the secretary must mark Official, Second dog, FEO, or BIS. Competitors do not choose Official, Second dog, or BIS on the public form.",
    "Drag a run to its required position inside the applicable class/order. Save the order before leaving the page.",
    "Keep second dogs at the end of the component so the judge can change hides. Keep BIS after the second-dog group. Keep FEO after Regular and before BIS.",
    "If a judge approves a trial-day change in component order, Started Container still must be held first when offered.",
    "Re-check every component because the rulebook warns that a competitor’s position may change between components.",
]: bullet(x)
image("03-running-order-workbook.png","Figure 8. Header from an actual TrialDesk running-order export. Each day receives its own worksheet with a day total, trial metadata, level-coloured class bands, component columns, reactive indicators, and class totals. Review the complete workbook’s names, streams, run groups, conflicts, and sequence before posting.",6.8)
heading("5.3 Export and post",2)
step(1,"Save the running order","Resolve any on-screen conflict warnings.")
step(2,"Export running order XLSX","Open it in Excel and inspect every day/level/component.")
step(3,"Print or post copies","Place a current copy at check-in and/or the handler staging area. Mark superseded copies clearly.")

page_break(); heading("6. Component-specific move-ups",1)
para("A qualifying component permits that same component at the next level: Started → Advanced or Advanced → Excellent. An Excellent title across all three components is required for Elite, so TrialDesk does not treat an individual Excellent component pass as an automatic Elite move-up.")
heading("6.1 When to approve a move-up",2)
for x in [
    "The dog has a verified qualifying result in the same lower-level component.",
    "The destination level/component is offered on the relevant trial day.",
    "The move-up does not duplicate the same component in the same trial.",
    "The stream remains correct: after a processed title, continued entries at the titled level must be Working; at the next level the dog may again compete as Amateur unless the handler/dog otherwise belongs in Working.",
    "The secretary has evidence even if the qualification occurred at an earlier trial. Record the source in the activity/audit context.",
]: bullet(x)
heading("6.2 Apply the move-up in TrialDesk",2)
step(1,"Open Running orders","Choose the day containing the run to be changed.")
step(2,"Find the component run","Confirm dog, handler, day, stream, and current level.")
step(3,"Select Move up","Review the confirmation showing the old and new level.")
step(4,"Confirm and regenerate","TrialDesk records the previous level, approval time, and approving user, then places the run in the destination offering.")
step(5,"Reorganize both affected orders","A move-up changes class counts and may create handler or official conflicts. Drag, save, re-export, and repost the affected order.")
step(6,"Undo only for a genuine correction","Use Undo move-up to return the run to its recorded prior level; the change remains auditable.")
callout("Do not infer titles", "Title Watch is an aid. A component move-up requires verified eligibility; a championship or Gold calculation must rely on official SDDA records, not competitor self-reporting.","warning")

page_break(); heading("7. Generate official score sheets",1)
image("15-score-sheets.png","Figure 9. Official Score Sheets. Packet controls remain disabled until accepted runs have a saved running position.",5.6)
heading("7.1 Prerequisites",2)
for x in ["Entries are accepted and unwanted entries are waitlisted/withdrawn/cancelled","Running order and all approved move-ups are saved","Trial number, date, judge, dog call name, breed, SDDA number, stream, and formal alerts are correct","The printer can print the official page size at 100% without fit-to-page distortion"]: bullet(x)
heading("7.2 Export Scent sheets",2)
step(1,"Open Score sheets","Confirm the displayed count against the accepted Scent run count.")
step(2,"Choose the required packet","Export the prefilled official PDF for the selected trial/day or offered set.")
step(3,"Inspect representative pages","Open at least one page for every offered level/component combination—not merely one page per level.")
step(4,"Verify mapped fields","Check trial number/date, call name, breed, SDDA number, formal alerts, stream checkbox, FEO marker, and footer sequence.")
step(5,"Print working and spare copies","Provide originals to the judge and copies for any shadow judge according to SDDA requirements.")
image("05-score-sheet-example.png","Figure 10. Mapping quality-control example. Every level and component uses an independent official template and coordinate map; never assume one sheet’s coordinates apply to another.",4.7)
heading("7.3 Export Games sheets",2)
para("Use the Games packet for Aerial, Distance, Speed, and Team. Verify the separate Game mapping, FEO mark, Aerial High/Highfly selection, and Team partner block. Games are Pass/Not Pass rather than numerically scored.")
callout("Critical print check", "If any text crosses a form line, checkbox X is outside its box, formal alert is missing, or the packet count differs from accepted runs, stop and correct the data or mapping before trial day.","warning")

page_break(); heading("8. Run the trial: secretary desk sequence",1)
heading("8.1 Before competitors arrive",2)
for x in [
    "Print/post the latest running orders and label each revision.",
    "Sort score sheets in run order by day, level, component/stream, and group.",
    "Prepare check-in list, registration-pending list, payment list, volunteer assignments, incident materials, ribbons, and spare supplies.",
    "Confirm search-area order with the judge. Started Container, if offered, is first and indoors.",
    "Keep clean scent supplies separate from previously scented items; hides are the judge’s responsibility unless arrangements were made.",
]: bullet(x)
heading("8.2 Check-in",2)
step(1,"Confirm identity","Verify handler, dog call name, SDDA registration number/card, and any pending correction.")
step(2,"Confirm operational notes","Ask about BIS, reactive-dog handling, accommodations, withdrawals, volunteer duties, second dogs, and Team partner changes.")
step(3,"Confirm order","Tell the competitor that position may differ by component and direct them to the posted order.")
step(4,"Record changes immediately","Update the entry/run group/order in TrialDesk and replace affected printed postings.")
heading("8.3 During each component",2)
for x in [
    "Officials run before duties when their duties can be temporarily reassigned.",
    "Regular entries run before second dogs; the judge changes hide location(s) for the second-dog group.",
    "FEO runs after Regular entries and before BIS. BIS runs last and must follow the current SDDA equipment requirements.",
    "The steward calls teams from a staging area without visual access to the search area.",
    "The judge completes score and time, signs the sheet, and returns it to the secretary. The secretary may total the sheet.",
    "All teams complete the component before score sheets are released/presented.",
]: bullet(x)
heading("8.4 Between components or days",2)
para("Record results, identify verified move-ups, change the next running order, re-run conflict checks, save, export, and repost. Never rely on handwritten changes alone when the next score packet or official workbook will be generated from TrialDesk.")

page_break(); heading("9. Enter scores and corrections",1)
image("16-score-entry.png","Figure 11. Audited Score Entry. Filter cards, use Save and next while working, and use the sticky Save all changes button before taking a break.",6.8)
heading("9.1 Scent scoring",2)
step(1,"Open Score entry","Only accepted runs appear.")
step(2,"Select the exact run","Match day, level, component, stream, dog, and handler to the signed sheet.")
step(3,"Enter result","Choose Qualifying, Non-qualifying, Absent, Withdrawn, or Excused as appropriate.")
step(4,"Enter score and time","For Qualifying/NQ results, transcribe the signed totals and time. Started components total 30/40/30; Advanced, Excellent, and Elite total 60/80/60.")
step(5,"Save","Every initial save and correction is recorded in the activity journal.")
heading("9.2 Games scoring",2)
para("Select Pass, Fail, Absent, Withdrawn, or Excused and enter the official time where applicable. Games are not numerically scored. Match Team results to the finalized pair and Aerial results to the correct High/Highfly category.")
heading("9.3 Correcting a score",2)
for x in ["Keep the signed sheet as the authority","Open the same run, amend only the incorrect fields, and save","Confirm that Results, Title Watch, Finances (if status changed), and the workbook now reflect the correction","Use Activity to identify who changed the result and compare before/after values"]: bullet(x)
callout("Paper control", "Retain and distribute score-sheet copies according to current SDDA instructions. TrialDesk’s digital audit supports—not replaces—the signed official sheet.")

page_break(); heading("10. Placements, results, and titles",1)
image("17-results.png","Figure 12. Results and placements. TrialDesk calculates provisional placements from recorded qualifying results and excludes FEO from placements.",6.8)
heading("10.1 Provisional placements",2)
para("Open Results after scoring. TrialDesk separates Scent by day, level, component, and stream. It includes qualifying runs only, ranks highest score first, then faster time, and excludes FEO. Exact score/time ties share a place. Games placements include passing Regular entries and rank faster time first; FEO is excluded.")
step(1,"Check completion badge","Resolve missing scores before treating placements as final.")
step(2,"Compare with signed sheets","Spot-check leaders, ties, NQs/fails, absences, withdrawals, and FEO exclusions.")
step(3,"Print / Save PDF","Use the results page print control for posting or presentation.")
heading("10.2 Title Watch",2)
image("18-title-ribbons.png","Figure 13. Close to Titles and Ribbon Planning. The top summary gives the maximum possible ribbon count, with Working reviews and title details below.",4.5)
para("Title Watch combines current accepted/scored entries with the latest public SDDA dog-history workbook snapshot. It highlights historical component Qs, possible same-trial Special titles, processed titles requiring Working at that level, and component-specific opportunities.")
for x in [
    "A standard Started/Advanced/Excellent title requires a qualifying score in Container, Interior, and Exterior at that level.",
    "An untitled dog passing all three components of a level at the same trial may earn the Special designation.",
    "Elite requires all three components passed at the same trial.",
    "A Games title requires three passes in that Game; Games Championship requires titles in at least three Games.",
    "Gold/championship progress is not accepted from competitor self-reporting. Confirm it against detailed official SDDA records.",
]: bullet(x)
callout("Authority", "Title Watch is planning support, not SDDA title issuance. If a dog is unmatched, history is incomplete, or a title affects stream eligibility, verify with SDDA before representing the status as official.","warning")

page_break(); heading("11. Finances, workbook, backup, and closeout",1)
heading("11.1 Financial Summary: start here",2)
para("Open the trial, then Financial Summary in the left menu. Amounts are Canadian dollars. Accepted entries generate charges automatically from Scent and Games pricing; do not enter the same fees again as a manual charge.")
image("06-finances-overview.png","Figure 14. Live Example Trial financial summary, September 2, 2026. Figures demonstrate the screen layout, not recommended entry prices. Judge and SDDA estimates are separate from actual expenses.",6.3)
para("Gross charges include adjustments. Fees waived are not money collected. Collected is payments minus refunds; Outstanding and Credits show handler balances. Cash net subtracts actual expenses from net collections, without deducting the displayed estimates again.")

page_break(); heading("11.1a Record a handler payment",2)
step(1,"Find the handler","Scroll to Handler accounts - payments and waivers. Search by handler or dog; use Owing, Credit, Settled, or All balances as needed.")
step(2,"Review the dogs","Matching handler name and email group dogs together. Entries without email remain separate. Check the dogs' charges and balances before recording money.")
step(3,"Select Record payment","For multiple dogs, keep All dogs - allocate payment, or choose one dog. The amount starts with the current balance; replace it with the amount actually received.")
image("07-record-payment.png","Figure 15. Payment received dialog. This unsaved example starts at $200.00 for two dogs. No money is processed by this form; it records money already received.",6.2)
step(4,"Enter payment details and Save","Confirm date, payment method, reference, and notes. Save once, wait for the dialog to close, and check the updated net paid and balance.")
step(5,"Check partial or excess payments","A partial payment leaves an outstanding balance. An all-dogs payment fills positive dog balances in order; excess remains as credit on the last dog. Inspect each dog when another dog already has credit.")

page_break(); heading("11.1b Waivers, restoration, refunds, and corrections",2)
step(1,"Waive an unpaid fee","On the specific dog's row, choose Waive fees. Enter all or part of the current outstanding balance and explain why in Notes. Save and check Fees waived and the balance; collected money does not increase.")
image("08-waive-fees.png","Figure 16. Fees waived dialog. The reason is required; a waiver cannot exceed this dog's outstanding balance. The example was opened and cancelled without saving.",6.0)
step(2,"Restore a waiver","Locate the waiver in Ledger - transaction history and choose Restore fees. Add a reason and save. It restores the original amount once and retains both records; waivers cannot be edited or deleted.")
step(3,"Record money returned","Choose Refund for that dog, enter the amount, date, method, reference and reason, then Save. Refund is unavailable when no net payment exists, and cannot exceed that dog's net payments.")
para("A refund returns recorded money but does not cancel the entry charge. If the fee should no longer be owed, separately correct the entry or use an appropriate waiver/credit adjustment. A waiver is a fixed amount, not free future entries: recheck balances after adding runs or changing prices.")
para("Use Adjust for an explained additional charge (positive) or credit (negative). In the ledger, use Edit to correct an existing transaction; it cannot switch dogs or transaction type. Use Delete only for a genuinely mistaken ordinary record. Corrections remain audited; do not clear the activity journal to fix money records.")

page_break(); heading("11.1c Expenses and financial statements",2)
step(1,"Select Add expense","Choose Trial expense, Judge expense, Volunteer expense, or SDDA remittance paid. Enter the actual amount, date, method, reference, payee, and a useful description; then Save.")
image("09-add-expense.png","Figure 17. Add expense dialog, opened without saving. Record an actual cost once; an estimate is not a payment.",6.0)
step(2,"Reconcile the ledger","Compare receipts, cash and transfers with recorded payments, refunds and expenses. Review each handler's Payment and waiver history. Judge estimates do not include all travel, meals or accommodation costs.")
step(3,"Export Excel","The sticky header exports Entry balances, Ledger, Summary, and Judge estimates worksheets. Entry balances contains numeric amounts and a balance formula; retain the file with the trial records.")
step(4,"Print statement","Use Print statement in a browser that supports printing and inspect print preview before printing or saving as PDF. Clear account search/filter controls first for a complete on-screen account list.")
callout("Preserve the record", "Back up the trial before major corrections. Use Activity Journal to review who recorded or changed information. Completed trials must be reopened before financial changes; repeat statements and backups after corrections.")

page_break(); heading("11.2 Official SDDA workbook",2)
image("19-official-workbook.png","Figure 18. Official Workbook. Export is enabled only when required trial identity and score checks pass; open the result in Excel to recalculate the official formulas.",6.8)
step(1,"Finish accepted-entry scoring","Resolve registration numbers, trial numbers, judges, results, scores, and times.")
step(2,"Open Official workbook","Read blockers and warnings. Correct source data rather than editing around a blocker.")
step(3,"Export each day group","The official workbook has two day sections. TrialDesk groups a 3–4 day event into separate one/two-day workbook exports.")
step(4,"Open in Excel","Allow original formulas to recalculate. Review Trial Info, level sheets, Games, Summary, High-in-Trial, fees, labels, formatted results, and the embedded SDDA Dogs registry.")
step(5,"Submit outside TrialDesk","TrialDesk downloads a completed copy; it does not upload directly to SDDA. Follow SDDA’s current submission/email instructions.")
heading("11.3 Backup and closeout",2)
image("21-trial-closeout.png","Figure 19. Trial Closeout. Resolve required items, review final exports, download the JSON backup, and only then complete and lock the trial.",6.8)
step(1,"Open Trial closeout","Resolve red blockers and review amber warnings.")
step(2,"Download complete trial backup","Store the JSON backup with the final workbook, results, running order, financial statement, and retained score-sheet records.")
step(3,"Complete and lock","Completion prevents operational edits. Use Reopen only for an authorized correction and repeat exports afterward.")
callout("Archive", "Keep a durable copy outside the browser. A downloaded backup is only useful if it is stored in a known, backed-up location and can be associated with the correct trial number/date.")
heading("11.4 Activity Journal",2)
image("20-activity-journal.png","Figure 20. Activity Journal. Search or filter the permanent secretary-facing history and print it for the closeout package.",5.6)
para("The journal groups trial setup, pricing, offering changes, CSV import batches, entry changes, payments, waivers, score corrections, and completion. It identifies who performed the action and shows the meaningful change without unnecessary database fields.")
step(1,"Search or filter","Use the search box for an action, person, dog, or record and the record-type list to narrow the journal.")
step(2,"Review material changes","Confirm imports are grouped as batches and corrections name the affected entry, score, payment, or trial setting.")
step(3,"Print journal","Save a PDF or printed copy with the final trial records before locking the trial.")

page_break(); heading("12. Troubleshooting and control checklists",1)
heading("12.1 Common problems",2)
table([
    ["Symptom", "Secretary action"],
    ["First login fails; Retry works", "Synchronize Windows date/time and time zone; close tabs and sign in again. The production login should use a full authenticated navigation."],
    ["Trial or Finances unavailable", "Confirm the signed-in profile has access to the trial and that the required migrations were applied."],
    ["Entry absent from order/sheets", "Confirm the entry is Accepted, not merely Received; verify at least one valid run and offering."],
    ["CSV row rejected", "Read the row-specific error; correct header mapping, identity, day, level, component, stream, or duplicate data."],
    ["FEO option missing", "Enable Allow FEO entries for that exact offering and save setup. Regular remains the default."],
    ["Move-up unavailable", "Confirm lower-level component qualification, destination offering, and permitted next level. Elite requires the Excellent title."],
    ["Wrong sheet count", "Compare accepted operational runs; exclude waitlisted/withdrawn/cancelled entries and check every imported selection."],
    ["PDF fields misaligned", "Stop printing and report the exact level/component/Game. Each official template has its own mapping."],
    ["Workbook blocker", "Correct the underlying trial/entry/score data, then export a fresh untouched official workbook copy."],
], widths=[2.05,4.75])
heading("12.2 Final pre-trial checklist",2)
for x in ["All intended entries accepted and balances reviewed","Registration-pending list resolved or actively controlled","Every run assigned to the correct day/level/component/stream/group","Officials, second dogs, FEO, and BIS sequenced correctly","Conflict checks reviewed after every move-up and manual reorder","One page from every score-sheet template visually checked","Running order exported, printed, and posted","Desk supplies, spare sheets, timers, ribbons, and backup equipment ready"]: bullet("☐ "+x)
page_break(); heading("12.3 Final closeout checklist",2)
for x in ["All signed sheets returned and results entered","Placements reviewed; FEO excluded","Title opportunities reviewed against official history","Payments and expenses reconciled","Official workbook exported, recalculated in Excel, and reviewed","Results, workbook, finances, activity, and running order archived","Complete JSON backup downloaded","Trial completed and locked after all required submissions are ready"]: bullet("☐ "+x)

page_break(); heading("13. Page and button glossary",1)
para("Use this section as a quick lookup when you know the control name but not where it is located. Controls that change data are disabled when prerequisites are incomplete or the trial is completed and locked.")
heading("13.1 Dashboard and trial setup",2)
table([
    ["Button or control", "Location", "Function"],
    ["Create trial / Create New Trial", "Dashboard or sidebar", "Opens the new-trial form."],
    ["Create SDDA Trial", "All Trials", "Opens the new-trial form."],
    ["Open trial", "All Trials card", "Opens the selected trial workspace."],
    ["Edit trial", "All Trials card", "Opens identity, format, and guarded trial-day editing for the selected trial."],
    ["Copy trial", "All Trials card or Edit trial", "Starts a clean draft using reusable setup from the selected trial."],
    ["Delete draft", "Draft trial card", "Permanently removes an unused draft after confirmation; unavailable after operational use."],
    ["Open setup", "Dashboard", "Opens Trial Details for the selected trial."],
    ["All trials / View all", "Dashboard or sidebar", "Opens the complete trial list."],
    ["Built-in entry form", "Dashboard", "Opens the selected trial's competitor form."],
    ["Google Form CSV / Entry roster", "Dashboard", "Opens Entries for import or review."],
    ["Add trial day", "Create trial", "Adds another date, up to the application limit of four."],
    ["Remove day", "Create trial", "Removes that date before the draft is created."],
    ["Cancel", "Create trial", "Returns to Trials without creating a draft."],
    ["Create Draft Trial", "Create trial", "Creates the editable trial workspace; it does not open entries."],
    ["Save trial changes", "Edit trial", "Saves the trial name, host club, venue/address, and format without changing entries or runs."],
    ["Add trial day", "Edit trial", "Adds another empty trial day, up to four."],
    ["Remove Day", "Edit trial", "Removes only the final day and is blocked while that day has offerings or runs."],
    ["Save trial days", "Edit trial", "Atomically saves chronological date changes while preserving existing day assignments."],
    ["Create draft copy", "Copy trial", "Creates a new draft with reusable setup but no operational or financial records."],
    ["Save setup", "Trial Details sticky header", "Saves offering, Games, and setup changes currently pending."],
    ["Save public details", "Trial Details", "Saves secretary contact, payment, and cancellation text."],
    ["Select all / Clear Scent", "Trial offering setup", "Selects or clears all Scent level/component offerings."],
    ["Select all / Clear Games", "Trial offering setup", "Selects or clears all Games offerings."],
    ["Save Scent pricing", "Scent entry pricing", "Saves component, package, and Elite fees."],
    ["Open entries and enable form", "Trial operations", "Changes the trial to Entries Open and activates submissions."],
    ["Close entries", "Trial operations", "Stops new public submissions for the whole trial."],
    ["Open or Close this day", "Trial day header", "Allows or stops entries for only that day."],
    ["Save day details", "Trial day", "Saves the date, SDDA trial number, and day judge."],
    ["Offering tile", "Trial day offerings", "Turns that Scent component or Game on or off."],
    ["Allow FEO entries", "Selected offering", "Shows or removes FEO for that exact public offering."],
    ["Refresh from SDDA", "SDDA Dog Registry", "Stages and activates the newest official SDDA Dogs list; administrator only."],
    ["Search registry", "SDDA Dog Registry", "Looks up an SDDA dog number or call name in the active registry."],
], widths=[1.9,1.8,3.9])

heading("13.2 Navigation sharing and team access",2)
table([
    ["Button or control", "Location", "Function"],
    ["Trial workflow stage", "Top of Trial Details", "Opens the named stage and shows whether its prerequisite is ready."],
    ["Competitor Entry Form", "Sidebar or Trial operations", "Opens the public form in a separate page."],
    ["Copy Entry Link", "Sidebar or Trial operations", "Copies the public form address to the clipboard."],
    ["Trial Team", "Sidebar or Trial operations", "Opens access management for the trial."],
    ["Add to trial", "Trial Team", "Adds a registered user or updates an existing member's role."],
    ["Trash icon", "Trial Team member", "Removes a non-owner member after confirmation."],
    ["Close to Titles and Ribbons", "Dashboard, Trial Details, or sidebar", "Opens title opportunities, Working review, and maximum ribbon planning."],
    ["Sign In", "Login", "Authenticates the secretary and opens the dashboard."],
    ["Forgot your password", "Login", "Requests a Supabase password-reset email."],
    ["Register here", "Login", "Opens account registration and email confirmation."],
], widths=[1.9,1.8,3.9])

heading("13.3 Competitor entry form",2)
table([
    ["Button or control", "Location", "Function"],
    ["Load my entry", "Top of public form", "Loads an editable entry when registration/confirmation number and email both match."],
    ["Check SDDA registry", "Competitor and dog step", "Checks the entered dog number and call name against the active registry."],
    ["Back / Continue", "Entry form steps", "Moves between Competitor and dog, Runs, Policies, and Review without submitting."],
    ["Submit entry", "Review step", "Creates the entry after required data, selections, declaration, and policies pass validation."],
    ["Save entry changes", "Review step while editing", "Saves permitted competitor or secretary corrections."],
    ["Print receipt", "Entry received page", "Prints the received selections, estimated amount, and return instructions."],
    ["Download receipt PDF", "Entry received page", "Downloads a formatted receipt containing the active private edit link."],
    ["Edit entry", "Entry received page", "Returns to step one with the submitted information populated."],
    ["Copy private entry link", "Entry received page", "Copies the confidential edit link; share only with that competitor."],
], widths=[1.9,1.8,3.9])

heading("13.4 Entries and import",2)
table([
    ["Button or control", "Location", "Function"],
    ["Search", "Entries", "Finds a handler, dog, or SDDA number."],
    ["Status filter", "Entries", "Shows All, Received, Accepted, Waitlisted, Rejected, or other statuses."],
    ["Export mailing list XLSX", "Entries", "Exports email, handler, received selections, and amount owing."],
    ["Import Google Form CSV", "Entries", "Selects and previews a familiar Google response CSV."],
    ["Import entries", "CSV import review", "Creates valid entries and reports rejected rows as one audited batch."],
    ["Edit entry", "Entry card", "Opens the secretary edit form with the current entry repopulated."],
    ["Secretary decision", "Entry card", "Changes Received to Accepted, Waitlisted, Rejected, Withdrawn, or another permitted status."],
    ["Print or Save PDF", "Entry receipt/review", "Prints the entry receipt with selections, amount, and edit link."],
], widths=[1.9,1.8,3.9])

heading("13.5 Running order score sheets and scoring",2)
table([
    ["Button or control", "Location", "Function"],
    ["Day Level Component filters", "Running Orders", "Select the exact order being reviewed."],
    ["Apply SDDA order", "Running Orders", "Applies Officials, Regular, Second dog, FEO, and BIS grouping rules."],
    ["Run group list", "Running Orders row", "Marks the run's operational group when the secretary must change it."],
    ["Move up", "Running Orders row", "Applies a component-specific approved move-up to an offered next level."],
    ["Drag handle / arrows", "Running Orders row", "Moves a run manually; re-check conflicts afterward."],
    ["Save", "Running Orders", "Persists the displayed order and group changes."],
    ["Export running order XLSX", "Running Orders", "Downloads the formatted multi-day running-order workbook."],
    ["Export complete packet", "Score Sheets", "Downloads all eligible Scent sheets as one PDF."],
    ["Export day packet", "Score Sheets", "Downloads one day's eligible Scent sheets."],
    ["Export Games packet", "Score Sheets", "Downloads all or one day's Games sheets."],
    ["Scent or Games filter", "Score Entry", "Switches scoring discipline and shows its recorded count."],
    ["Trial day filter / Search", "Score Entry", "Narrows cards without changing scores."],
    ["Save and next", "Score card", "Saves that run and advances to the next displayed card."],
    ["Save all changes", "Score Entry sticky header", "Saves every changed score card before a break."],
], widths=[1.9,1.8,3.9])

heading("13.6 Results titles finances and closeout",2)
table([
    ["Button or control", "Location", "Function"],
    ["Print or Save PDF", "Results", "Prints provisional placements after score review."],
    ["Export Excel", "Close to Titles", "Exports title opportunities, ribbon totals, and Working reviews."],
    ["Print or Save PDF", "Close to Titles", "Prints the title and ribbon-planning report."],
    ["Refresh history", "Close to Titles", "Refreshes official SDDA workbook history used for title proximity."],
    ["Open Entry Roster", "Close to Titles", "Opens entries requiring Working review."],
    ["Print statement / Export Excel", "Finances", "Produces the current financial statement or workbook."],
    ["Add expense", "Finances", "Records an operating, judge, volunteer, or SDDA remittance expense."],
    ["Record payment", "Handler account", "Allocates money already received across selected dogs."],
    ["Waive fees", "Dog balance", "Waives an amount with a required reason; does not remove runs."],
    ["Refund", "Dog balance", "Records money returned against collected funds."],
    ["Adjust", "Dog balance", "Records a documented financial correction."],
    ["Restore waiver", "Payment and waiver history", "Reverses a prior waiver when permitted."],
    ["Pencil / Trash", "Financial ledger", "Edits or deletes a ledger item with confirmation and audit history."],
    ["Export official workbook", "Official Workbook", "Fills a fresh official workbook after blockers pass."],
    ["Print journal", "Activity Journal", "Prints or saves the filtered secretary-facing audit record."],
    ["Results Workbook Finances Activity links", "Trial Closeout", "Opens each final export for review."],
    ["Download complete trial backup", "Trial Closeout", "Downloads the complete JSON recovery record."],
    ["Complete and lock trial", "Trial Closeout", "Locks operational editing after required checks pass."],
    ["Reopen for corrections", "Completed trial closeout", "Reopens a completed trial and records the action."],
], widths=[1.9,1.8,3.9])

page_break(); heading("Sources and scope notes",1)
p=doc.add_paragraph(); hyperlink(p,"SDDA Master Rule Book v5.1 (effective July 1, 2026)","https://www.sdda.ca/wp-content/uploads/2026/06/SDDA-Rules-July-2026.pdf")
p=doc.add_paragraph(); hyperlink(p,"SDDA Rules and Rulebook page","https://www.sdda.ca/front-page/rules-rulebook/")
p=doc.add_paragraph(); hyperlink(p,"SDDA website","https://www.sdda.ca/")
para("The operating sequence, primary pages, and visible controls were reviewed against K:\\TrialManager through September 14, 2026. Screenshots show read-only inspection and previously prepared unsaved finance dialogs; no live transactions or operational changes were submitted for this manual. The current SDDA rulebook and official forms remain controlling.")
callout("Maintenance", "Revise this manual whenever SDDA publishes a new rulebook, official score sheet, Games form, results workbook, fee schedule, or submission instruction—or whenever TrialDesk changes a workflow shown here.")

# Keep figures inline and add basic accessibility descriptions to drawings.
for drawing, desc in zip(doc.element.body.iter(qn("w:drawing")), [
    "TrialDesk dashboard", "Create SDDA Trial form", "Trial setup workflow and checklist",
    "Trial Team access management", "Entry roster controls", "Google Form response spreadsheet",
    "Running order controls", "Running order workbook", "Official score sheet packet page",
    "Official SDDA score sheet mapping examples", "Audited score entry controls",
    "Provisional results and placements", "Close to Titles and ribbon planning",
    "Financial summary and judge estimates", "Record a handler payment",
    "Waive a dog's outstanding fees", "Record an operating expense",
    "Official SDDA workbook export page", "Trial closeout checklist", "Activity Journal"
]):
    docPr = drawing.find(".//" + qn("wp:docPr"))
    if docPr is not None: docPr.set("descr", desc)

doc.core_properties.title = "SDDA TrialDesk Secretary Manual"
doc.core_properties.subject = "Step-by-step SDDA trial administration workflow"
doc.core_properties.author = "SDDA TrialDesk"
doc.core_properties.keywords = "SDDA, TrialDesk, secretary, running order, score sheets, move-ups, placements"
doc.save(OUT)
print(OUT)
